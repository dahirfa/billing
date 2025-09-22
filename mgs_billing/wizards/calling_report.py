from odoo import models, fields, api
from odoo.exceptions import ValidationError
import xlsxwriter
import base64
import io
from openpyxl import Workbook
import logging
_logger = logging.getLogger(__name__)

class CallingReport(models.TransientModel):
    _name = 'mgs_billing.calling_report_wiz'
    _description = 'Calling Report'

    date_from = fields.Date(
        default=fields.Date.today().replace(day=1), required=True)
    date_to = fields.Date(default=fields.Date.today(), required=True)

    zone_id = fields.Many2one('mgs_billing.zone')
    collector_id = fields.Many2one('res.partner', string='Collector', domain=[('is_collector', '=', True)])
    states = fields.Selection(
        [('all', 'All'), ('posted', 'Posted')], default="posted", string='Target Moves', required=True)
    company_id = fields.Many2one(
        'res.company', string='Company', default=lambda self: self.env.user.company_id.id)
    greater_less = fields.Selection(
        [('Greater', 'Greater than'), ('Less', 'Less than')], default="Greater", string='Show', required=True)
    greater_less_amount = fields.Float(
        string='Amount', default=1, required=True)
    datas = fields.Binary('File', readonly=True)
    datas_fname = fields.Char('Filename', readonly=True)
    property_state = fields.Selection(
        [("connected", "Connected"), ("disconnected", "Disconnected")], default="connected", required=True)

    @api.constrains('date_from', 'date_to')
    def _check_the_date_from_and_to(self):
        if self.date_to and self.date_from and self.date_to < self.date_from:
            raise ValidationError('''From Date should be less than To Date.''')

    def confirm(self):
        data = {
            'ids': self.ids,
            'model': self._name,
            'form': {
                'date_from': self.date_from,
                'date_to': self.date_to,
                'property_state': self.property_state,
                'zone_id': [self.zone_id.id, self.zone_id.name],
                'collector_id': [self.collector_id.id, self.collector_id.name],
                'states': self.states,
                'company_id': [self.company_id.id, self.company_id.name],
                'greater_less': self.greater_less,
                'greater_less_amount': self.greater_less_amount,
            },
        }
                
        _logger.info(data)
        return self.env.ref('mgs_billing.action_calling_report').report_action(self, data=data)


    def export_to_excel(self):
        # params = [
        #     self.date_from,  
        #     self.date_from,  
        #     self.date_to,    
        #     self.date_to,    
        #     self.date_to,
        # ]
        # query = """
        #     SELECT
        #         rp.property_id,
        #         mbp.name AS property_name,
        #         rp.name AS partner_name,
        #         rp.phone AS sender_phone,
        #         rp.mobile AS tenant_mobile,
        #         COALESCE(SUM(CASE WHEN ai.date < %s THEN ai.amount_total - ai.amount_residual ELSE 0 END), 0) AS initial_balance,
        #         COALESCE(SUM(CASE WHEN ai.date >= %s AND ai.date <= %s THEN ai.amount_total ELSE 0 END), 0) AS invoiced,
        #         COALESCE(SUM(CASE WHEN ai.date <= %s THEN ai.amount_total - ai.amount_residual ELSE 0 END), 0) AS balance
        #     FROM account_move ai
        #     JOIN res_partner rp ON ai.partner_id = rp.id
        #     JOIN mgs_billing_property mbp ON rp.property_id = mbp.id
        #     WHERE ai.move_type = 'out_invoice'
        #     AND ai.state = 'posted'
        #     AND ai.date <= %s
        #     GROUP BY rp.property_id, mbp.name, rp.name, rp.phone, rp.mobile
        #     ORDER BY mbp.name;
        # """

        # self.env.cr.execute(query, tuple(params))
        # rows = self.env.cr.dictfetchall()
        
        params = [self.date_from, self.date_from, self.date_to,
                  self.date_from, self.date_to, self.date_to]
        query = """
        SELECT
            row_number() over (order by rp.id DESC) as id,
            rp.id AS billing_account_id,
            rp.name AS partner_name,
            rp.company_id AS company_id,
            mbz.name AS zone_name,
            collector.id AS collector_id,
            collector.name AS collector_name,
            mbp.name as property_name,
            rp.mobile as tenant_mobile,
            rp.alternative_number as sender_phone,
            
            COALESCE(sum(CASE WHEN aml.date < %s THEN aml.debit-aml.credit else 0.0 END), 0) AS initial_balance,
            COALESCE(sum(CASE WHEN aml.date between %s and %s THEN aml.debit else 0.0 END), 0) invoiced,
            COALESCE(sum(CASE WHEN aml.date between %s  and %s  THEN aml.credit else 0.0 END), 0) paid,
            COALESCE(sum (aml.debit-aml.credit), 0.0) balance
        FROM res_partner rp
            LEFT JOIN mgs_billing_property mbp ON rp.property_id = mbp.id
            LEFT JOIN mgs_billing_zone mbz ON mbp.zone_id = mbz.id
            LEFT JOIN res_partner collector ON mbz.collector_id=collector.id
            left join res_users as ru on ru.partner_id=collector.id
            LEFT JOIN account_move_line aml ON aml.partner_id=rp.id
            LEFT JOIN account_account AS aa ON aml.account_id = aa.id
        WHERE aa.account_type = 'asset_receivable' and aml.partner_id IS NOT NULL AND aml.date <= %s
            AND aml.parent_state = 'posted' AND rp.is_shareholder != True
            AND rp.is_tenancy = True
        """

        if self.zone_id:
            params.append(self.zone_id.id)
            query += " AND mbz.id = %s"
            
        if self.property_state:
            params.append(self.property_state)
            query += " AND mbp.state = %s"

        if self.collector_id:
            params.append(self.collector_id.id)
            query += " and collector.id = %s"

        if self.company_id:
            params.append(self.company_id.id)
            query += " and aml.company_id = %s"

        query += """
        GROUP BY rp.id, rp.company_id, mbz.id, collector.id, mbp.name, rp.mobile, rp.name"""
        query += " HAVING COALESCE(sum (aml.debit-aml.credit), 0.0) " + \
            "> " + str(self.greater_less_amount) if self.greater_less == 'Greater' else "< " + \
            str(self.greater_less_amount)

        self.env.cr.execute(query, tuple(params))
        rows = self.env.cr.dictfetchall()
        
        _logger.info("=================== CHECK =====================")
        _logger.info(rows)
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Calling Report"

        # ==== HEADER SECTION ====
        row_num = 1
        ws.cell(row=row_num, column=1, value="Calling Report")
        ws.cell(row=row_num, column=1).font = ws.cell(row=row_num, column=1).font.copy(bold=True, size=14)

        row_num += 2
        ws.cell(row=row_num, column=1, value="Date From:")
        ws.cell(row=row_num, column=2, value=str(self.date_from))
        row_num += 1
        ws.cell(row=row_num, column=1, value="Date To:")
        ws.cell(row=row_num, column=2, value=str(self.date_to))

        if self.zone_id:
            row_num += 1
            ws.cell(row=row_num, column=1, value="Zone:")
            ws.cell(row=row_num, column=2, value=self.zone_id.name)

        if self.collector_id:
            row_num += 1
            ws.cell(row=row_num, column=1, value="Collector:")
            ws.cell(row=row_num, column=2, value=self.collector_id.name)

        if self.property_state:
            row_num += 1
            ws.cell(row=row_num, column=1, value="Property State:")
            ws.cell(row=row_num, column=2, 
                    value=dict(self._fields['property_state'].selection).get(self.property_state))

        row_num += 2  # leave a blank line before table

        # ==== TABLE HEADER ====
        headers = [
            "No", "Property ID", "Billing Account", "Sender",
            "Tenant Phone", "Prev Bal", "Curr Bill", "Curr Balance"
        ]
        ws.append(headers)

        # ==== TABLE DATA ====
        counter = 1
        for line in rows:
            ws.append([
                counter,
                line.get('property_name'),
                line.get('partner_name'),
                line.get('sender_phone'),
                line.get('tenant_mobile'),
                line.get('initial_balance'),
                line.get('invoiced'),
                line.get('balance'),
            ])
            counter += 1

        # Save to memory
        fp = io.BytesIO()
        wb.save(fp)
        fp.seek(0)
        file_data = base64.b64encode(fp.read())

        # Store in wizard
        self.write({
            'datas': file_data,
            'datas_fname': "Calling Report.xlsx",
        })

        # Download
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/?model=%s&id=%s&field=datas&filename_field=datas_fname&download=true' % (
                self._name, self.id),
            'target': 'self',
        }




